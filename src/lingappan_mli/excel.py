"""Formatted Excel export for Lingappan MLI results."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F5F9F")
ID_HEADER_FILL = PatternFill("solid", fgColor="334155")
MLI_HEADER_FILL = PatternFill("solid", fgColor="166534")
NON_AIRSPACE_HEADER_FILL = PatternFill("solid", fgColor="6D28D9")
AIRSPACE_HEADER_FILL = PatternFill("solid", fgColor="0F766E")
QC_HEADER_FILL = PatternFill("solid", fgColor="C2410C")
HEADER_FONT = Font(color="FFFFFF", bold=True)
THIN_BORDER = Border(bottom=Side(style="thin", color="D0D5DD"))
AUTOSIZE_SAMPLE_ROWS = 200
NUMBER_FORMAT_MAX_ROWS = 20_000


NUMERIC_FORMATS = {
    "um": "0.00",
    "um2": "0.00",
    "fraction": "0.0000",
    "area": "0.00",
    "count": "0",
    "pixels": "0.00",
    "px": "0.00",
}


def _format_for_header(header: str) -> str | None:
    lower = header.lower()
    if lower in {
        "area_measurement_method",
        "area_measurement_connectivity",
        "area_column_semantics",
        "area_measurement_method_description",
    }:
        return None
    if lower.endswith("_per_field") or "per_field" in lower:
        return "0.00"
    if (
        lower.endswith("_count")
        or lower.endswith("count")
        or "_count_" in lower
        or "boundary_intersections" in lower
        or lower.endswith("_chords")
        or "_chords_" in lower
        or lower.endswith("_seed")
        or lower in {"width_px", "height_px"}
    ):
        return "0"
    if "fraction" in lower or "ratio" in lower:
        return "0.0000"
    if "um2" in lower or "area" in lower:
        return "0.00"
    if "um" in lower or lower.endswith("_px") or "pixels" in lower:
        return "0.00"
    return None


def _autosize_columns(ws) -> None:
    sample_rows = min(ws.max_row, AUTOSIZE_SAMPLE_ROWS)
    for column_cells in ws.iter_cols(max_row=sample_rows):
        max_length = 0
        col_idx = column_cells[0].column
        column_letter = get_column_letter(col_idx)
        for cell in column_cells:
            if cell.value is None:
                continue
            max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[column_letter].width = min(max(max_length + 2, 10), 48)


def _header_fill(header: str) -> PatternFill:
    lower = header.lower()
    if lower in {"file", "slide", "field", "fields", "measurement", "orientation", "line"}:
        return ID_HEADER_FILL
    if "non-airspace" in lower:
        return NON_AIRSPACE_HEADER_FILL
    if "mli" in lower:
        return MLI_HEADER_FILL
    if "airspace" in lower:
        return AIRSPACE_HEADER_FILL
    if any(token in lower for token in ("warning", "status", "reference", "qc")):
        return QC_HEADER_FILL
    return HEADER_FILL


def _style_sheet(ws) -> None:
    if ws.max_row < 1 or ws.max_column < 1:
        return
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        header = str(cell.value) if cell.value is not None else ""
        cell.fill = _header_fill(header)
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
    ws.row_dimensions[1].height = 28

    headers = [str(cell.value) if cell.value is not None else "" for cell in ws[1]]
    format_end_row = min(ws.max_row, NUMBER_FORMAT_MAX_ROWS)
    for col_idx, header in enumerate(headers, start=1):
        number_format = _format_for_header(header)
        if number_format is None:
            continue
        for row in range(2, format_end_row + 1):
            ws.cell(row=row, column=col_idx).number_format = number_format

    alignment_end_row = min(ws.max_row, AUTOSIZE_SAMPLE_ROWS)
    for row in ws.iter_rows(min_row=2, max_row=alignment_end_row):
        for cell in row:
            cell.alignment = Alignment(vertical="top")

    _autosize_columns(ws)


def write_results_workbook(path: str | Path, sheets: dict[str, pd.DataFrame]) -> None:
    """Write and style an Excel workbook.

    Parameters
    ----------
    path:
        Output `.xlsx` path.
    sheets:
        Mapping of worksheet names to DataFrames.
    """
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        for name, frame in sheets.items():
            safe_name = name[:31]
            frame.to_excel(writer, sheet_name=safe_name, index=False)

    wb = load_workbook(out)
    wb.properties.title = "Lingappan MLI Analyzer Results"
    wb.properties.subject = "Mean linear intercept morphometry"
    wb.properties.creator = "Lingappan MLI Analyzer"
    for ws in wb.worksheets:
        _style_sheet(ws)
    wb.save(out)
