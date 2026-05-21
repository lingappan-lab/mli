from __future__ import annotations

import json

import numpy as np
import pandas as pd
from openpyxl import load_workbook
import pytest
from PIL import Image

import lingappan_mli.analysis as analysis_module
from lingappan_mli.analysis import (
    DEFAULT_CALIBRATION_SOURCE,
    DEFAULT_PIXEL_HEIGHT_UM,
    DEFAULT_PIXEL_WIDTH_UM,
    UNRECORDED_CALIBRATION_SOURCE,
    AnalysisParams,
    analyze_image,
    process_files,
    summarize_by_slide,
)
from lingappan_mli.contours import (
    AREA_COLUMN_SEMANTICS,
    AREA_MEASUREMENT_METHOD,
    airspace_component_connectivity_label,
    analyze_airspace_contours,
)
from lingappan_mli.measure import (
    _runs,
    measure_phase_chords,
    summarize_measurements,
    summarize_phase_line_intercepts,
)
from lingappan_mli.grid import GridLine, generate_grid, generate_grid_layout
from lingappan_mli.io import load_image_rgb, load_image_rgb_with_metadata
from lingappan_mli.thresholding import threshold_airspace


def _line_positions(lines, orientation: str) -> list[int]:
    return [line.position for line in lines if line.orientation == orientation]


_CHORD_SORT_COLUMNS = [
    "measure_type",
    "phase",
    "orientation",
    "line_id",
    "start_y",
    "start_x",
    "end_y",
    "end_x",
]


def _sorted_chord_frame(chords: pd.DataFrame) -> pd.DataFrame:
    sort_columns = [column for column in _CHORD_SORT_COLUMNS if column in chords.columns]
    if not sort_columns:
        return chords.reset_index(drop=True)
    return chords.sort_values(sort_columns).reset_index(drop=True)


def _measure_direct_phase(
    mask: np.ndarray,
    lines: list[GridLine],
    *,
    phase: str,
    pixel_width_um: float = 1.0,
    pixel_height_um: float = 1.0,
    exclude_edge_touching: bool = True,
    min_chord_um: float = 0.0,
) -> pd.DataFrame:
    return measure_phase_chords(
        mask,
        lines,
        phase=phase,
        pixel_width_um=pixel_width_um,
        pixel_height_um=pixel_height_um,
        exclude_edge_touching=exclude_edge_touching,
        min_chord_um=min_chord_um,
    )


def _rgb_image(shape: tuple[int, int], *bright_regions, background=(40, 40, 40), bright=(240, 240, 240)):
    arr = np.zeros((*shape, 3), dtype=np.uint8)
    arr[:, :] = background
    for region in bright_regions:
        arr[region] = bright
    return arr


def test_grid_random_offsets_are_optional_seeded_and_reported_for_count_and_spacing():
    default_count = generate_grid((100, 80), orientation="both", strategy="count", num_lines=4)
    random_count_a = generate_grid(
        (100, 80),
        orientation="both",
        strategy="count",
        num_lines=4,
        random_offset=True,
        random_seed=123,
    )
    random_count_b = generate_grid(
        (100, 80),
        orientation="both",
        strategy="count",
        num_lines=4,
        random_offset=True,
        random_seed=123,
    )
    random_count_c = generate_grid(
        (100, 80),
        orientation="both",
        strategy="count",
        num_lines=4,
        random_offset=True,
        random_seed=456,
    )

    assert _line_positions(default_count, "Horizontal") == [20, 40, 59, 79]
    assert _line_positions(random_count_a, "Horizontal") == _line_positions(random_count_b, "Horizontal")
    assert _line_positions(random_count_a, "Vertical") == _line_positions(random_count_b, "Vertical")
    assert _line_positions(random_count_a, "Horizontal") != _line_positions(default_count, "Horizontal")
    assert _line_positions(random_count_a, "Horizontal") != _line_positions(random_count_c, "Horizontal")

    default_spacing = generate_grid(
        (100, 80),
        orientation="both",
        strategy="spacing",
        line_spacing_um=10,
        pixel_width_um=1,
        pixel_height_um=1,
    )
    random_spacing_a = generate_grid(
        (100, 80),
        orientation="both",
        strategy="spacing",
        line_spacing_um=10,
        pixel_width_um=1,
        pixel_height_um=1,
        random_offset=True,
        random_seed=123,
    )
    random_spacing_b = generate_grid(
        (100, 80),
        orientation="both",
        strategy="spacing",
        line_spacing_um=10,
        pixel_width_um=1,
        pixel_height_um=1,
        random_offset=True,
        random_seed=123,
    )
    assert _line_positions(default_spacing, "Horizontal") == [
        5,
        15,
        25,
        35,
        45,
        55,
        65,
        75,
        85,
        95,
    ]
    assert _line_positions(random_spacing_a, "Horizontal") == _line_positions(random_spacing_b, "Horizontal")
    assert _line_positions(random_spacing_a, "Vertical") == _line_positions(random_spacing_b, "Vertical")
    assert _line_positions(random_spacing_a, "Horizontal") != _line_positions(default_spacing, "Horizontal")

    layout = generate_grid_layout(
        (100, 80),
        orientation="horizontal",
        random_offset=True,
        random_seed=123,
    )
    assert layout.random_offset is True
    assert layout.random_seed == 123
    assert layout.horizontal_offset_px is not None
    assert layout.vertical_offset_px is None


def test_grid_random_seed_is_generated_when_offsets_enabled_without_seed():
    params = AnalysisParams(grid_random_offset=True).validated()
    assert params.grid_random_offset is True
    assert params.grid_random_seed is not None
    assert params.grid_random_seed >= 0


def test_calibration_metadata_records_physical_dimensions_and_source(tmp_path):
    arr = _rgb_image((4, 6), np.s_[1:3, 2:5], background=(30, 30, 30))
    img_path = tmp_path / "SlideCal_0001.png"
    Image.fromarray(arr).save(img_path)

    result = analyze_image(
        img_path,
        AnalysisParams(
            pixel_width_um=2.0,
            pixel_height_um=3.0,
            calibration_source="stage micrometer, 20x objective",
            threshold_method="otsu",
        ),
    )

    assert result.summary["calibration_source"] == "stage micrometer, 20x objective"
    assert result.summary["calibration_is_default"] is False
    assert result.summary["calibration_warning"] == ""
    assert np.isclose(result.summary["pixel_area_um2"], 6.0)
    assert np.isclose(result.summary["field_width_um"], 12.0)
    assert np.isclose(result.summary["field_height_um"], 12.0)
    assert np.isclose(result.summary["field_area_um2"], 144.0)

    manual_without_source = AnalysisParams(pixel_width_um=2.0, pixel_height_um=3.0).validated()
    assert manual_without_source.calibration_source == UNRECORDED_CALIBRATION_SOURCE


def test_default_calibration_is_warned_and_logged(tmp_path):
    arr = _rgb_image((4, 6), np.s_[1:3, 2:5], background=(30, 30, 30))
    img_path = tmp_path / "SlideDefaultCal_0001.png"
    Image.fromarray(arr).save(img_path)

    out = tmp_path / "default_calibration"
    result = process_files([img_path], out, AnalysisParams(threshold_method="otsu"))
    field_row = result["field_summary"].iloc[0]
    log_df = result["processing_log"]

    assert field_row["calibration_source"] == DEFAULT_CALIBRATION_SOURCE
    assert bool(field_row["calibration_is_default"]) is True
    assert "Default calibration is in use" in field_row["calibration_warning"]
    assert np.isclose(field_row["pixel_area_um2"], DEFAULT_PIXEL_WIDTH_UM * DEFAULT_PIXEL_HEIGHT_UM)
    assert np.isclose(field_row["field_width_um"], 6 * DEFAULT_PIXEL_WIDTH_UM)
    assert np.isclose(field_row["field_height_um"], 4 * DEFAULT_PIXEL_HEIGHT_UM)

    warning_rows = log_df[log_df["status"] == "warning"]
    assert len(warning_rows) == 1
    assert warning_rows.iloc[0]["filename"] == "calibration"
    assert "Default calibration is in use" in warning_rows.iloc[0]["message"]
    assert bool(warning_rows.iloc[0]["calibration_is_default"]) is True

    with open(out / "parameters.json", encoding="utf-8") as handle:
        parameters = json.load(handle)
    assert parameters["calibration_source"] == DEFAULT_CALIBRATION_SOURCE
    assert parameters["calibration_is_default"] is True
    assert "Default calibration is in use" in parameters["calibration_warning"]
    assert np.isclose(parameters["pixel_area_um2"], DEFAULT_PIXEL_WIDTH_UM * DEFAULT_PIXEL_HEIGHT_UM)

    with open(out / "README_results.md", encoding="utf-8") as handle:
        readme = handle.read()
    assert "Default-calibration warning" in readme


def test_non_uint8_image_scaling_is_recorded_and_logged(tmp_path):
    arr = np.array(
        [
            [100, 100, 250, 250],
            [100, 500, 750, 250],
            [100, 500, 1000, 250],
            [100, 100, 250, 250],
        ],
        dtype=np.uint16,
    )
    img_path = tmp_path / "Slide16_0001.tif"
    Image.fromarray(arr).save(img_path)

    loaded = load_image_rgb_with_metadata(img_path)
    assert loaded.rgb.dtype == np.uint8
    assert loaded.rgb.shape == (4, 4, 3)
    assert int(loaded.rgb[0, 0, 0]) == 0
    assert int(loaded.rgb[2, 2, 0]) == 255
    assert loaded.metadata.original_dtype == "uint16"
    assert loaded.metadata.scaling_applied is True
    assert loaded.metadata.scaling_min == 100.0
    assert loaded.metadata.scaling_max == 1000.0
    assert "min-max scaled to uint8" in " | ".join(loaded.metadata.warnings)

    result = process_files(
        [img_path],
        tmp_path / "non_uint8",
        AnalysisParams(pixel_width_um=1, pixel_height_um=1, threshold_method="otsu", num_lines=2),
    )
    field_row = result["field_summary"].iloc[0]
    assert field_row["image_original_dtype"] == "uint16"
    assert bool(field_row["image_scaling_applied"]) is True
    assert field_row["image_scaling_min"] == 100.0
    assert field_row["image_scaling_max"] == 1000.0
    assert "min-max scaled to uint8" in field_row["image_load_warnings"]

    warning_rows = result["processing_log"][result["processing_log"]["status"] == "warning"]
    assert len(warning_rows) == 1
    assert warning_rows.iloc[0]["filename"] == img_path.name
    assert "min-max scaled to uint8" in warning_rows.iloc[0]["message"]
    assert bool(warning_rows.iloc[0]["image_scaling_applied"]) is True


def test_multiframe_tiff_first_frame_use_is_recorded_and_logged(tmp_path):
    frame0 = np.zeros((5, 5), dtype=np.uint8)
    frame0[0, 0] = 10
    frame0[2, 2] = 220
    frame1 = np.full((5, 5), 240, dtype=np.uint8)
    img_path = tmp_path / "SlidePages_0001.tif"
    Image.fromarray(frame0).save(img_path, save_all=True, append_images=[Image.fromarray(frame1)])

    loaded = load_image_rgb_with_metadata(img_path)
    assert loaded.metadata.original_format == "TIFF"
    assert loaded.metadata.frame_index == 0
    assert loaded.metadata.page_count == 2
    assert int(loaded.rgb[0, 0, 0]) == 10
    assert int(loaded.rgb[2, 2, 0]) == 220
    assert "only frame/page 0" in " | ".join(loaded.metadata.warnings)

    result = process_files(
        [img_path],
        tmp_path / "multiframe",
        AnalysisParams(pixel_width_um=1, pixel_height_um=1, threshold_method="otsu", num_lines=2),
    )
    field_row = result["field_summary"].iloc[0]
    assert field_row["image_original_format"] == "TIFF"
    assert field_row["image_frame_index"] == 0
    assert field_row["image_page_count"] == 2
    assert bool(field_row["image_scaling_applied"]) is False
    assert "only frame/page 0" in field_row["image_load_warnings"]

    warning_rows = result["processing_log"][result["processing_log"]["status"] == "warning"]
    assert len(warning_rows) == 1
    assert warning_rows.iloc[0]["filename"] == img_path.name
    assert "only frame/page 0" in warning_rows.iloc[0]["message"]
    assert warning_rows.iloc[0]["image_page_count"] == 2


def test_embedded_resolution_metadata_is_recorded_and_mismatch_is_logged(tmp_path):
    arr = np.zeros((6, 6, 3), dtype=np.uint8)
    arr[:, :] = [40, 40, 40]
    arr[1:5, 1:5] = [240, 240, 240]
    img_path = tmp_path / "SlideDPI_0001.tif"
    Image.fromarray(arr).save(img_path, dpi=(25400, 12700))

    loaded = load_image_rgb_with_metadata(img_path)
    assert loaded.metadata.embedded_resolution_source == "TIFF resolution tags"
    assert loaded.metadata.embedded_resolution_unit == "inch"
    assert loaded.metadata.embedded_resolution_x == 25400.0
    assert loaded.metadata.embedded_resolution_y == 12700.0
    assert loaded.metadata.embedded_pixel_width_um == 1.0
    assert loaded.metadata.embedded_pixel_height_um == 2.0
    assert loaded.metadata.warnings == ()

    result = process_files(
        [img_path],
        tmp_path / "embedded_resolution",
        AnalysisParams(
            pixel_width_um=2,
            pixel_height_um=2,
            calibration_source="stage micrometer",
            threshold_method="otsu",
            num_lines=2,
        ),
    )
    field_row = result["field_summary"].iloc[0]
    assert field_row["image_embedded_resolution_source"] == "TIFF resolution tags"
    assert field_row["image_embedded_resolution_unit"] == "inch"
    assert field_row["image_embedded_pixel_width_um"] == 1.0
    assert field_row["image_embedded_pixel_height_um"] == 2.0
    assert "Embedded image resolution metadata" in field_row["image_load_warnings"]

    warning_rows = result["processing_log"][result["processing_log"]["status"] == "warning"]
    assert len(warning_rows) == 1
    assert warning_rows.iloc[0]["filename"] == img_path.name
    assert "Embedded image resolution metadata" in warning_rows.iloc[0]["message"]
    assert warning_rows.iloc[0]["image_embedded_pixel_width_um"] == 1.0


def test_batch_field_size_consistency_warnings_are_logged(tmp_path):
    first_image = tmp_path / "SlideBatch_0001.png"
    second_image = tmp_path / "SlideBatch_0002.png"
    Image.fromarray(_rgb_image((8, 8), (slice(2, 6), slice(2, 6)))).save(first_image)
    Image.fromarray(_rgb_image((8, 12), (slice(2, 6), slice(2, 10)))).save(second_image)

    result = process_files(
        [first_image, second_image],
        tmp_path / "batch_qc",
        AnalysisParams(
            pixel_width_um=1,
            pixel_height_um=1,
            calibration_source="stage micrometer",
            threshold_method="otsu",
            num_lines=2,
        ),
    )

    warning_rows = result["processing_log"][result["processing_log"]["status"] == "warning"]
    messages = " | ".join(warning_rows["message"].astype(str))
    assert set(warning_rows["filename"]) == {"batch_qc"}
    assert "multiple analyzed image pixel dimensions" in messages
    assert "Batch physical field size varies" in messages


def test_process_files_preflight_rejects_oversized_images(tmp_path):
    image_path = tmp_path / "SlideHuge_0001.png"
    Image.fromarray(_rgb_image((8, 8), (slice(2, 6), slice(2, 6)))).save(image_path)

    with pytest.raises(ValueError, match="Image preflight failed") as exc_info:
        process_files(
            [image_path],
            tmp_path / "preflight",
            AnalysisParams(pixel_width_um=1, pixel_height_um=1),
            max_image_pixels=10,
        )

    assert "per-image limit" in str(exc_info.value)


def test_multiframe_tiff_exif_orientation_is_applied_after_frame_metadata(tmp_path):
    base_frame = np.array([[1, 2, 3], [4, 5, 6]], dtype=np.uint8)
    frame0 = np.stack([base_frame] * 3, axis=-1)
    frame1 = np.full((2, 3, 3), 9, dtype=np.uint8)
    exif = Image.Exif()
    exif[274] = 6  # rotate 90 degrees clockwise
    img_path = tmp_path / "SlideOriented_0001.tif"
    Image.fromarray(frame0).save(
        img_path,
        save_all=True,
        append_images=[Image.fromarray(frame1)],
        exif=exif,
    )

    loaded = load_image_rgb_with_metadata(img_path)
    assert loaded.metadata.page_count == 2
    assert loaded.rgb.shape == (3, 2, 3)
    assert loaded.rgb[:, :, 0].tolist() == [[4, 1], [5, 2], [6, 3]]
    assert load_image_rgb(img_path).shape == (3, 2, 3)


def test_runs_reports_signed_transitions_and_edge_touching_runs():
    assert _runs(np.array([], dtype=bool)) == []
    assert _runs(np.array([True, True, False, True, False, False, True, True], dtype=bool)) == [
        (0, 1, True),
        (3, 3, False),
        (6, 7, True),
    ]
    assert _runs(np.array([0, 255, 255, 0, 1, 0], dtype=np.uint8)) == [
        (1, 2, False),
        (4, 4, False),
    ]


def test_synthetic_all_airspace_and_all_tissue_masks_have_expected_edge_censored_chords():
    lines = [
        GridLine("H001", "Horizontal", 0, 1, 3, 1),
        GridLine("V001", "Vertical", 2, 0, 2, 2),
    ]
    pixel_width_um = 0.5
    pixel_height_um = 2.0

    all_airspace = np.ones((3, 4), dtype=bool)
    included_airspace = _measure_direct_phase(
        all_airspace,
        lines,
        phase="airspace",
        pixel_width_um=pixel_width_um,
        pixel_height_um=pixel_height_um,
        exclude_edge_touching=False,
    )
    edge_censored_airspace = _measure_direct_phase(
        all_airspace,
        lines,
        phase="airspace",
        pixel_width_um=pixel_width_um,
        pixel_height_um=pixel_height_um,
        exclude_edge_touching=True,
    )
    assert len(included_airspace) == 2
    assert included_airspace["edge_touching"].all()
    assert sorted(included_airspace["length_px"].astype(int).tolist()) == [3, 4]
    assert included_airspace.set_index("orientation")["length_um"].to_dict() == {
        "Horizontal": 2.0,
        "Vertical": 6.0,
    }
    assert edge_censored_airspace.empty
    assert _measure_direct_phase(all_airspace, lines, phase="tissue").empty

    all_airspace_area = analyze_airspace_contours(
        all_airspace,
        pixel_width_um=pixel_width_um,
        pixel_height_um=pixel_height_um,
        exclude_edge_touching=True,
    )
    assert all_airspace_area.metrics["airspace_component_count_all"] == 1
    assert all_airspace_area.metrics["final_airspace_area_pixels_mask"] == 0
    assert all_airspace_area.metrics["total_airspace_area_pixels_mask"] == 12
    assert np.isclose(all_airspace_area.metrics["largest_airspace_component_fraction_of_airspace"], 1.0)
    assert all_airspace_area.metrics["largest_airspace_component_touches_edge"] is True

    all_tissue = np.zeros((3, 4), dtype=bool)
    included_tissue = _measure_direct_phase(
        all_tissue,
        lines,
        phase="tissue",
        pixel_width_um=pixel_width_um,
        pixel_height_um=pixel_height_um,
        exclude_edge_touching=False,
    )
    edge_censored_tissue = _measure_direct_phase(
        all_tissue,
        lines,
        phase="tissue",
        pixel_width_um=pixel_width_um,
        pixel_height_um=pixel_height_um,
        exclude_edge_touching=True,
    )
    assert len(included_tissue) == 2
    assert included_tissue["edge_touching"].all()
    assert sorted(included_tissue["length_px"].astype(int).tolist()) == [3, 4]
    assert included_tissue.set_index("orientation")["length_um"].to_dict() == {
        "Horizontal": 2.0,
        "Vertical": 6.0,
    }
    assert edge_censored_tissue.empty
    assert _measure_direct_phase(all_tissue, lines, phase="airspace").empty

    all_tissue_area = analyze_airspace_contours(
        all_tissue,
        pixel_width_um=pixel_width_um,
        pixel_height_um=pixel_height_um,
        exclude_edge_touching=True,
    )
    assert all_tissue_area.metrics["airspace_component_count_all"] == 0
    assert all_tissue_area.metrics["total_airspace_area_pixels_mask"] == 0
    assert all_tissue_area.metrics["final_airspace_area_pixels_mask"] == 0
    assert all_tissue_area.metrics["largest_airspace_component_touches_edge"] is None


def test_synthetic_striped_mask_min_chord_boundary_direct_overlay_and_indirect_formula():
    mask = np.array([[False, True, True, False, True, True, True, False]], dtype=bool)
    lines = [GridLine("H001", "Horizontal", 0, 0, 7, 0)]

    boundary_inclusive = _measure_direct_phase(
        mask,
        lines,
        phase="airspace",
        pixel_width_um=0.75,
        pixel_height_um=2.0,
        exclude_edge_touching=True,
        min_chord_um=1.5,
    )
    assert boundary_inclusive["length_px"].astype(int).tolist() == [2, 3]
    assert np.allclose(boundary_inclusive["length_um"].to_numpy(), [1.5, 2.25])
    assert not boundary_inclusive["edge_touching"].any()

    just_above_boundary = _measure_direct_phase(
        mask,
        lines,
        phase="airspace",
        pixel_width_um=0.75,
        pixel_height_um=2.0,
        exclude_edge_touching=True,
        min_chord_um=1.5001,
    )
    assert just_above_boundary["length_px"].astype(int).tolist() == [3]
    assert np.isclose(just_above_boundary.iloc[0]["length_um"], 2.25)

    summary = summarize_phase_line_intercepts(
        mask,
        lines,
        phase="airspace",
        pixel_width_um=0.75,
        pixel_height_um=2.0,
        exclude_edge_touching=True,
        min_chord_um=1.5,
        measure_type="MLI",
    )
    assert np.isclose(summary["mli_total_airspace_line_length_um_accepted"], 3.75)
    assert summary["mli_airspace_boundary_intersection_count_accepted"] == 4
    assert summary["mli_chord_count_accepted"] == 2
    assert np.isclose(
        summary["mli_indirect_equivalent_mean_um"],
        2 * 3.75 / 4,
    )
    assert np.isclose(
        summary["mli_indirect_equivalent_mean_um"],
        boundary_inclusive["length_um"].mean(),
    )


def test_synthetic_diagonal_mask_connectivity_and_edge_censoring_are_explicit():
    mask = np.eye(3, dtype=bool)

    eight_connected = analyze_airspace_contours(
        mask,
        pixel_width_um=2.0,
        pixel_height_um=3.0,
        exclude_edge_touching=True,
        connectivity=8,
    )
    four_connected = analyze_airspace_contours(
        mask,
        pixel_width_um=2.0,
        pixel_height_um=3.0,
        exclude_edge_touching=True,
        connectivity=4,
    )

    assert eight_connected.metrics["airspace_component_count_all"] == 1
    assert eight_connected.metrics["airspace_component_count_edge_excluded"] == 1
    assert eight_connected.metrics["final_airspace_area_pixels_mask"] == 0
    assert np.isclose(eight_connected.metrics["total_airspace_area_um2_mask"], 18.0)

    assert four_connected.metrics["airspace_component_count_all"] == 3
    assert four_connected.metrics["airspace_component_count_edge_excluded"] == 2
    assert four_connected.metrics["final_airspace_area_pixels_mask"] == 1
    assert np.isclose(four_connected.metrics["final_airspace_area_um2_mask"], 6.0)


def test_synthetic_threshold_boundary_uses_strictly_greater_for_bright_airspace():
    image = np.array([[10, 10, 200, 200]], dtype=np.uint8)

    gray, bright_airspace, threshold = threshold_airspace(
        image,
        method="otsu",
        airspace_bright=True,
    )
    assert gray.tolist() == [[10, 10, 200, 200]]
    assert threshold == 10
    assert bright_airspace.tolist() == [[False, False, True, True]]

    _, dark_airspace, dark_threshold = threshold_airspace(
        image,
        method="otsu",
        airspace_bright=False,
    )
    assert dark_threshold == threshold
    assert dark_airspace.tolist() == [[True, True, False, False]]


def test_chord_measurement_simple_horizontal_airspace():
    mask = np.zeros((5, 10), dtype=bool)
    mask[2, 2:7] = True
    lines = generate_grid(mask.shape, orientation="horizontal", strategy="count", num_lines=1)
    df = measure_phase_chords(
        mask,
        lines,
        phase="airspace",
        pixel_width_um=1.0,
        pixel_height_um=1.0,
    )
    assert len(df) == 1
    assert int(df.iloc[0]["length_px"]) == 5


def test_direct_measurement_handles_mli_non_airspace_both_orientations():
    mask = np.array(
        [
            [False, False, True, True, False, False, True, False, False],
            [False, True, True, True, False, True, True, True, False],
            [False, True, False, True, True, True, False, True, False],
            [True, True, False, False, True, True, True, False, False],
            [False, False, True, False, False, True, True, True, True],
            [False, True, True, False, True, False, False, True, False],
            [True, False, True, False, False, True, True, False, True],
        ],
        dtype=bool,
    )
    lines = [
        GridLine("H001", "Horizontal", 0, 1, 8, 1),
        GridLine("H002", "Horizontal", 0, 3, 8, 3),
        GridLine("V001", "Vertical", 2, 0, 2, 6),
        GridLine("V002", "Vertical", 6, 0, 6, 6),
    ]

    for phase, measure_type in (("airspace", "MLI"), ("non_airspace", "Non-airspace")):
        chords = _measure_direct_phase(
            mask,
            lines,
            phase=phase,
            pixel_width_um=0.5,
            pixel_height_um=2.0,
            exclude_edge_touching=True,
        )
        assert set(chords["measure_type"]) == {measure_type}
        assert set(chords["orientation"]) == {"Horizontal", "Vertical"}
        if phase == "airspace":
            horizontal_lengths = chords[chords["orientation"] == "Horizontal"]["length_um"]
            vertical_lengths = chords[chords["orientation"] == "Vertical"]["length_um"]
            assert np.isclose(horizontal_lengths.iloc[0], 1.5)
            assert np.isclose(vertical_lengths.iloc[0], 4.0)



def test_direct_measurement_applies_min_chord_and_edge_filters():
    mask = np.array(
        [
            [False, False, True, True, False, False, True, False, False],
            [False, True, True, True, False, True, True, True, False],
            [False, True, False, True, True, True, False, True, False],
            [True, True, False, False, True, True, True, False, False],
            [False, False, True, False, False, True, True, True, True],
            [False, True, True, False, True, False, False, True, False],
            [True, False, True, False, False, True, True, False, True],
        ],
        dtype=bool,
    )
    lines = [
        GridLine("H001", "Horizontal", 0, 1, 8, 1),
        GridLine("H002", "Horizontal", 0, 3, 8, 3),
        GridLine("V001", "Vertical", 2, 0, 2, 6),
        GridLine("V002", "Vertical", 6, 0, 6, 6),
    ]

    expected_counts = {"airspace": (8, 4), "tissue": (3, 2)}
    for phase, (include_edge_count, accepted_count) in expected_counts.items():
        include_edge = _measure_direct_phase(
            mask,
            lines,
            phase=phase,
            pixel_width_um=1.0,
            pixel_height_um=1.0,
            exclude_edge_touching=False,
            min_chord_um=2.0,
        )
        accepted = _measure_direct_phase(
            mask,
            lines,
            phase=phase,
            pixel_width_um=1.0,
            pixel_height_um=1.0,
            exclude_edge_touching=True,
            min_chord_um=2.0,
        )
        assert len(include_edge) == include_edge_count
        assert len(accepted) == accepted_count
        assert set(include_edge["orientation"]) == {"Horizontal", "Vertical"}
        assert set(accepted["orientation"]) == {"Horizontal", "Vertical"}
        assert include_edge["edge_touching"].any()
        assert not accepted["edge_touching"].any()
        assert (include_edge["length_um"] >= 2.0).all()
        assert (accepted["length_um"] >= 2.0).all()


def test_direct_mli_aliases_and_indirect_equivalent_metrics_match_complete_chords():
    mask = np.array([[False, True, True, False, True, True, True, False]], dtype=bool)
    lines = generate_grid(mask.shape, orientation="horizontal", strategy="count", num_lines=1)
    chords = measure_phase_chords(
        mask,
        lines,
        phase="airspace",
        pixel_width_um=1.0,
        pixel_height_um=1.0,
        exclude_edge_touching=True,
    )

    summary = summarize_measurements(chords, "MLI")
    summary.update(
        summarize_phase_line_intercepts(
            mask,
            lines,
            phase="airspace",
            pixel_width_um=1.0,
            pixel_height_um=1.0,
            exclude_edge_touching=True,
            measure_type="MLI",
        )
    )

    assert summary["mli_chord_count"] == 2
    assert summary["mli_chord_count_accepted"] == 2
    assert summary["mli_chord_count_raw"] == 2
    assert np.isclose(summary["mli_mean_um"], 2.5)
    assert summary["mli_direct_mean_um"] == summary["mli_mean_um"]
    assert summary["mli_pooled_chord_mean_um"] == summary["mli_mean_um"]
    assert summary["mli_direct_orientation_balanced_mean_um"] == summary["mli_orientation_balanced_mean_um"]
    assert summary["mli_orientation_balanced_direct_mean_um"] == summary["mli_orientation_balanced_mean_um"]
    assert summary["mli_direct_horizontal_vertical_mean_ratio"] is None
    assert summary["mli_direct_horizontal_vertical_mean_delta_um"] is None
    assert np.isclose(summary["mli_total_airspace_line_length_um_raw"], 5.0)
    assert np.isclose(summary["mli_total_airspace_line_length_um_accepted"], 5.0)
    assert summary["mli_airspace_boundary_intersection_count_raw"] == 4
    assert summary["mli_airspace_boundary_intersection_count_accepted"] == 4
    assert np.isclose(summary["mli_indirect_equivalent_mean_um"], summary["mli_direct_mean_um"])
    assert np.isclose(
        summary["mli_indirect_equivalent_orientation_balanced_mean_um"],
        summary["mli_direct_orientation_balanced_mean_um"],
    )


def test_non_edge_airspace_component_area_metrics_and_censored_values():
    mask = np.zeros((5, 5), dtype=bool)
    mask[0, 0:2] = True  # edge-touching component; censored from non-edge area
    mask[2:4, 2] = True  # fully internal component

    edge_excluded = analyze_airspace_contours(
        mask,
        pixel_width_um=2.0,
        pixel_height_um=3.0,
        exclude_edge_touching=True,
    )
    metrics = edge_excluded.metrics
    assert metrics["area_measurement_method"] == AREA_MEASUREMENT_METHOD
    assert metrics["area_measurement_connectivity"] == airspace_component_connectivity_label()
    assert metrics["area_column_semantics"] == AREA_COLUMN_SEMANTICS
    assert metrics["total_airspace_area_pixels_mask"] == 4
    assert metrics["largest_airspace_component_area_pixels"] == 2
    assert np.isclose(metrics["largest_airspace_component_area_um2"], 12.0)
    assert np.isclose(metrics["largest_airspace_component_fraction_of_airspace"], 0.5)
    assert metrics["largest_airspace_component_touches_edge"] is True
    assert metrics["final_airspace_area_pixels_mask"] == 2
    assert (
        metrics["non_edge_airspace_component_area_pixels_mask"]
        == metrics["final_airspace_area_pixels_mask"]
    )
    assert np.isclose(metrics["non_edge_airspace_component_area_um2_mask"], 12.0)
    assert (
        metrics["non_edge_airspace_component_area_um2_mask"]
        == metrics["final_airspace_area_um2_mask"]
    )

    include_edges = analyze_airspace_contours(
        mask,
        pixel_width_um=2.0,
        pixel_height_um=3.0,
        exclude_edge_touching=False,
    )
    include_metrics = include_edges.metrics
    assert include_metrics["final_airspace_area_pixels_mask"] == 4
    assert include_metrics["non_edge_airspace_component_area_pixels_mask"] == 2


def test_airspace_component_connectivity_is_configurable_and_default_remains_8_connected():
    mask = np.zeros((4, 4), dtype=bool)
    mask[0, 0] = True  # edge pixel connected diagonally to internal pixels under 8-connectivity
    mask[1, 1] = True
    mask[2, 2] = True

    default_analysis = analyze_airspace_contours(mask, pixel_width_um=1.0, pixel_height_um=1.0)
    eight_connected = analyze_airspace_contours(mask, pixel_width_um=1.0, pixel_height_um=1.0, connectivity=8)
    four_connected = analyze_airspace_contours(mask, pixel_width_um=1.0, pixel_height_um=1.0, connectivity=4)

    assert default_analysis.metrics == eight_connected.metrics
    assert eight_connected.metrics["airspace_component_connectivity"] == 8
    assert eight_connected.metrics["area_measurement_connectivity"] == airspace_component_connectivity_label()
    assert eight_connected.metrics["airspace_component_count_all"] == 1
    assert eight_connected.metrics["airspace_component_count_edge_excluded"] == 1
    assert eight_connected.metrics["final_airspace_area_pixels_mask"] == 0

    assert four_connected.metrics["airspace_component_connectivity"] == 4
    assert four_connected.metrics["area_measurement_connectivity"] == airspace_component_connectivity_label(4)
    assert four_connected.metrics["airspace_component_count_all"] == 3
    assert four_connected.metrics["airspace_component_count_edge_excluded"] == 1
    assert four_connected.metrics["final_airspace_area_pixels_mask"] == 2
    assert four_connected.metrics["total_airspace_area_pixels_mask"] == 3

    with pytest.raises(ValueError, match="must be 4 or 8"):
        analyze_airspace_contours(mask, pixel_width_um=1.0, pixel_height_um=1.0, connectivity=6)
    with pytest.raises(ValueError, match="must be 4 or 8"):
        AnalysisParams(airspace_component_connectivity=6).validated()


def test_edge_exclusion_sensitivity_reports_raw_include_edge_excluded_and_accepted_for_mli_and_non_airspace():
    mask = np.array([[True, True, False, True, True, True, False, True]], dtype=bool)
    lines = generate_grid(mask.shape, orientation="horizontal", strategy="count", num_lines=1)

    mli_summary = summarize_phase_line_intercepts(
        mask,
        lines,
        phase="airspace",
        pixel_width_um=1.0,
        pixel_height_um=1.0,
        exclude_edge_touching=True,
        min_chord_um=2.0,
        measure_type="MLI",
    )
    assert mli_summary["mli_chord_count_raw"] == 3
    assert mli_summary["mli_chord_count_include_edge"] == 2
    assert mli_summary["mli_chord_count_min_length_excluded"] == 1
    assert mli_summary["mli_chord_count_edge_excluded"] == 1
    assert np.isclose(mli_summary["mli_chord_fraction_min_length_excluded"], 1 / 3)
    assert np.isclose(mli_summary["mli_chord_fraction_edge_excluded"], 0.5)
    assert mli_summary["mli_chord_count_accepted"] == 1
    assert np.isclose(mli_summary["mli_mean_um_raw"], 2.0)
    assert np.isclose(mli_summary["mli_mean_um_include_edge"], 2.5)
    assert np.isclose(mli_summary["mli_mean_um_edge_excluded"], 2.0)
    assert np.isclose(mli_summary["mli_mean_um_accepted"], 3.0)
    assert np.isclose(mli_summary["mli_orientation_balanced_mean_um_raw"], 2.0)
    assert np.isclose(mli_summary["mli_orientation_balanced_mean_um_include_edge"], 2.5)
    assert np.isclose(mli_summary["mli_orientation_balanced_mean_um_edge_excluded"], 2.0)
    assert np.isclose(mli_summary["mli_orientation_balanced_mean_um_accepted"], 3.0)

    airspace_for_non_airspace = np.array([[False, False, True, False, False, False, True, False]], dtype=bool)
    non_airspace_summary = summarize_phase_line_intercepts(
        airspace_for_non_airspace,
        lines,
        phase="non_airspace",
        pixel_width_um=1.0,
        pixel_height_um=1.0,
        exclude_edge_touching=True,
        min_chord_um=2.0,
        measure_type="Non-airspace",
    )
    assert non_airspace_summary["non_airspace_chord_count_raw"] == 3
    assert non_airspace_summary["non_airspace_chord_count_include_edge"] == 2
    assert non_airspace_summary["non_airspace_chord_count_min_length_excluded"] == 1
    assert non_airspace_summary["non_airspace_chord_count_edge_excluded"] == 1
    assert np.isclose(non_airspace_summary["non_airspace_chord_fraction_min_length_excluded"], 1 / 3)
    assert np.isclose(non_airspace_summary["non_airspace_chord_fraction_edge_excluded"], 0.5)
    assert non_airspace_summary["non_airspace_chord_count_accepted"] == 1
    assert np.isclose(non_airspace_summary["non_airspace_mean_um_raw"], 2.0)
    assert np.isclose(non_airspace_summary["non_airspace_mean_um_include_edge"], 2.5)
    assert np.isclose(non_airspace_summary["non_airspace_mean_um_edge_excluded"], 2.0)
    assert np.isclose(non_airspace_summary["non_airspace_mean_um_accepted"], 3.0)
    assert np.isclose(non_airspace_summary["non_airspace_orientation_balanced_mean_um_edge_excluded"], 2.0)


def test_qc_metrics_and_rejected_edge_overlays_are_exported(tmp_path):
    arr = _rgb_image((5, 8), background=(20, 20, 20))
    arr[2, 0:2] = [240, 240, 240]  # edge-touching, long enough to be edge-rejected
    arr[2, 3:6] = [240, 240, 240]  # accepted internal largest component
    arr[2, 7] = [240, 240, 240]  # short edge chord rejected by minimum length first
    img_path = tmp_path / "SlideQC_0001.png"
    Image.fromarray(arr).save(img_path)

    result = analyze_image(
        img_path,
        AnalysisParams(
            pixel_width_um=1,
            pixel_height_um=1,
            num_lines=1,
            orientation="horizontal",
            threshold_method="otsu",
            min_chord_um=2.0,
            exclude_edge_touching=True,
        ),
        tmp_path / "out",
    )

    summary = result.summary
    assert summary["largest_airspace_component_area_pixels"] == 3
    assert np.isclose(summary["largest_airspace_component_area_um2"], 3.0)
    assert np.isclose(summary["largest_airspace_component_fraction_of_airspace"], 0.5)
    assert summary["largest_airspace_component_touches_edge"] is False
    assert summary["mli_chord_count_raw"] == 3
    assert summary["mli_chord_count_include_edge"] == 2
    assert summary["mli_chord_count_min_length_excluded"] == 1
    assert summary["mli_chord_count_edge_excluded"] == 1
    assert summary["mli_chord_count_accepted"] == 1
    assert np.isclose(summary["mli_chord_fraction_min_length_excluded"], 1 / 3)
    assert np.isclose(summary["mli_chord_fraction_edge_excluded"], 0.5)

    image_dir = tmp_path / "out" / "SlideQC_0001"
    assert result.mask_overlay_path == image_dir / "06_qc" / "mask_overlay.png"
    assert result.mli_rejected_edge_overlay_path == image_dir / "06_qc" / "rejected_edge_chords_mli.png"
    assert result.mask_overlay_path.exists()
    assert result.mli_rejected_edge_overlay_path.exists()
    assert (image_dir / "03_mli" / "overlay_mli_combined.png").exists()
    assert (image_dir / "01_preprocessing" / "binary_airspace_mask.png").exists()
    rejected_overlay = np.asarray(Image.open(result.mli_rejected_edge_overlay_path).convert("RGB"))
    rejected_red_pixels = (
        (rejected_overlay[..., 0] > 200)
        & (rejected_overlay[..., 1] < 100)
        & (rejected_overlay[..., 2] < 100)
    )
    assert rejected_red_pixels.any()



def test_orientation_anisotropy_diagnostics_and_clear_mean_aliases_for_mli_and_non_airspace():
    chords = pd.DataFrame(
        [
            {"measure_type": "MLI", "orientation": "Horizontal", "length_um": 10.0},
            {"measure_type": "MLI", "orientation": "Horizontal", "length_um": 14.0},
            {"measure_type": "MLI", "orientation": "Vertical", "length_um": 3.0},
            {"measure_type": "Non-airspace", "orientation": "Horizontal", "length_um": 2.0},
            {"measure_type": "Non-airspace", "orientation": "Vertical", "length_um": 8.0},
            {"measure_type": "Non-airspace", "orientation": "Vertical", "length_um": 10.0},
        ]
    )

    mli_summary = summarize_measurements(chords, "MLI")
    assert np.isclose(mli_summary["mli_mean_um"], 9.0)
    assert np.isclose(mli_summary["mli_pooled_chord_mean_um"], mli_summary["mli_mean_um"])
    assert np.isclose(mli_summary["mli_orientation_balanced_mean_um"], 7.5)
    assert np.isclose(
        mli_summary["mli_orientation_balanced_direct_mean_um"],
        mli_summary["mli_orientation_balanced_mean_um"],
    )
    assert np.isclose(mli_summary["mli_direct_horizontal_vertical_mean_ratio"], 4.0)
    assert np.isclose(mli_summary["mli_direct_horizontal_vertical_mean_delta_um"], 9.0)

    non_airspace_summary = summarize_measurements(chords, "Non-airspace")
    assert np.isclose(non_airspace_summary["non_airspace_mean_um"], 20.0 / 3.0)
    assert np.isclose(
        non_airspace_summary["non_airspace_pooled_chord_mean_um"],
        non_airspace_summary["non_airspace_mean_um"],
    )
    assert np.isclose(non_airspace_summary["non_airspace_orientation_balanced_mean_um"], 5.5)
    assert np.isclose(
        non_airspace_summary["mean_non_airspace_chord_um"],
        non_airspace_summary["non_airspace_orientation_balanced_mean_um"],
    )
    assert np.isclose(
        non_airspace_summary["orientation_balanced_mean_non_airspace_chord_um"],
        non_airspace_summary["non_airspace_orientation_balanced_mean_um"],
    )
    assert np.isclose(
        non_airspace_summary["pooled_mean_non_airspace_chord_um"],
        non_airspace_summary["non_airspace_pooled_chord_mean_um"],
    )
    assert np.isclose(
        non_airspace_summary["non_airspace_orientation_balanced_direct_mean_um"],
        non_airspace_summary["non_airspace_orientation_balanced_mean_um"],
    )
    assert np.isclose(non_airspace_summary["non_airspace_direct_horizontal_vertical_mean_ratio"], 2.0 / 9.0)
    assert np.isclose(non_airspace_summary["non_airspace_direct_horizontal_vertical_mean_delta_um"], -7.0)

    horizontal_only = summarize_measurements(chords[chords["orientation"] == "Horizontal"], "MLI")
    assert horizontal_only["mli_direct_horizontal_vertical_mean_ratio"] is None
    assert horizontal_only["mli_direct_horizontal_vertical_mean_delta_um"] is None


def test_slide_summary_reports_field_balanced_chord_pooled_sem_and_chords_per_field():
    base = {
        "airspace_fraction": 0.5,
        "non_edge_airspace_fraction": 0.4,
        "total_airspace_area_um2_mask": 100.0,
        "final_airspace_area_um2_mask": 80.0,
        "non_edge_airspace_component_area_um2_mask": 80.0,
        "largest_airspace_component_area_um2": 30.0,
        "largest_airspace_component_fraction_of_airspace": 0.3,
    }
    rows = [
        {
            **base,
            "filename": "SlideWeight_0001.png",
            "slide_id": "SlideWeight",
            "mli_orientation_balanced_mean_um": 10.0,
            "mli_orientation_balanced_direct_mean_um": 10.0,
            "mli_direct_orientation_balanced_mean_um": 10.0,
            "mli_pooled_chord_mean_um": 5.0,
            "mli_chord_count": 2,
            "non_airspace_orientation_balanced_mean_um": 4.0,
            "non_airspace_orientation_balanced_direct_mean_um": 4.0,
            "non_airspace_pooled_chord_mean_um": 2.0,
            "non_airspace_chord_count": 1,
            "mean_non_airspace_chord_um": 4.0,
            "orientation_balanced_mean_non_airspace_chord_um": 4.0,
            "pooled_mean_non_airspace_chord_um": 2.0,
        },
        {
            **base,
            "filename": "SlideWeight_0002.png",
            "slide_id": "SlideWeight",
            "mli_orientation_balanced_mean_um": 30.0,
            "mli_orientation_balanced_direct_mean_um": 30.0,
            "mli_direct_orientation_balanced_mean_um": 30.0,
            "mli_pooled_chord_mean_um": 50.0,
            "mli_chord_count": 8,
            "non_airspace_orientation_balanced_mean_um": 10.0,
            "non_airspace_orientation_balanced_direct_mean_um": 10.0,
            "non_airspace_pooled_chord_mean_um": 20.0,
            "non_airspace_chord_count": 3,
            "mean_non_airspace_chord_um": 10.0,
            "orientation_balanced_mean_non_airspace_chord_um": 10.0,
            "pooled_mean_non_airspace_chord_um": 20.0,
        },
    ]

    slide_row = summarize_by_slide(pd.DataFrame(rows)).iloc[0]

    assert np.isclose(slide_row["mean_mli_um"], 20.0)
    assert np.isclose(slide_row["field_balanced_mean_mli_um"], 20.0)
    assert np.isclose(slide_row["sd_mli_um"], np.sqrt(200.0))
    assert np.isclose(slide_row["field_sd_mli_um"], np.sqrt(200.0))
    assert np.isclose(slide_row["field_sem_mli_um"], 10.0)
    assert np.isclose(slide_row["sem_mli_um"], slide_row["field_sem_mli_um"])
    assert np.isclose(slide_row["mean_mli_pooled_chord_um"], 27.5)
    assert np.isclose(slide_row["chord_pooled_mean_mli_um"], 41.0)
    assert np.isclose(slide_row["mean_mli_chords_per_field"], 5.0)
    assert slide_row["total_mli_chords"] == 10

    assert np.isclose(slide_row["mean_non_airspace_chord_um"], 7.0)
    assert np.isclose(slide_row["field_balanced_mean_non_airspace_chord_um"], 7.0)
    assert np.isclose(slide_row["field_sd_non_airspace_chord_um"], np.sqrt(18.0))
    assert np.isclose(slide_row["field_sem_non_airspace_chord_um"], 3.0)
    assert np.isclose(slide_row["sem_non_airspace_chord_um"], 3.0)
    assert np.isclose(slide_row["mean_pooled_non_airspace_chord_um"], 11.0)
    assert np.isclose(slide_row["chord_pooled_mean_non_airspace_chord_um"], 15.5)
    assert np.isclose(slide_row["mean_non_airspace_chords_per_field"], 2.0)
    assert slide_row["total_non_airspace_chords"] == 4
    assert np.isclose(slide_row["mean_non_edge_airspace_fraction"], 0.4)


def test_slide_summary_preserves_missing_variability_when_no_fields_are_measured():
    base = {
        "airspace_fraction": 0.5,
        "total_airspace_area_um2_mask": 100.0,
        "final_airspace_area_um2_mask": 80.0,
        "non_edge_airspace_component_area_um2_mask": 80.0,
        "largest_airspace_component_area_um2": 30.0,
        "largest_airspace_component_fraction_of_airspace": 0.3,
        "mli_orientation_balanced_mean_um": None,
        "mli_orientation_balanced_direct_mean_um": None,
        "mli_pooled_chord_mean_um": None,
        "mli_chord_count": 0,
        "non_airspace_orientation_balanced_mean_um": None,
        "non_airspace_orientation_balanced_direct_mean_um": None,
        "non_airspace_pooled_chord_mean_um": None,
        "non_airspace_chord_count": 0,
        "mean_non_airspace_chord_um": None,
        "orientation_balanced_mean_non_airspace_chord_um": None,
        "pooled_mean_non_airspace_chord_um": None,
    }
    rows = [
        {**base, "filename": "SlideMissing_0001.png", "slide_id": "SlideMissing"},
        {**base, "filename": "SlideMissing_0002.png", "slide_id": "SlideMissing"},
        {**base, "filename": "SlideMixed_0001.png", "slide_id": "SlideMixed"},
        {
            **base,
            "filename": "SlideMixed_0002.png",
            "slide_id": "SlideMixed",
            "mli_orientation_balanced_mean_um": 12.0,
            "mli_orientation_balanced_direct_mean_um": 12.0,
            "mli_pooled_chord_mean_um": 12.0,
            "mli_chord_count": 2,
            "non_airspace_orientation_balanced_mean_um": 5.0,
            "non_airspace_orientation_balanced_direct_mean_um": 5.0,
            "non_airspace_pooled_chord_mean_um": 5.0,
            "non_airspace_chord_count": 1,
            "mean_non_airspace_chord_um": 7.0,
            "orientation_balanced_mean_non_airspace_chord_um": 7.0,
            "pooled_mean_non_airspace_chord_um": 7.0,
        },
    ]

    slide_summary = summarize_by_slide(pd.DataFrame(rows)).set_index("slide_id")
    missing = slide_summary.loc["SlideMissing"]
    mixed = slide_summary.loc["SlideMixed"]

    assert missing["mli_measured_field_count"] == 0
    assert missing["non_airspace_measured_field_count"] == 0
    assert pd.isna(missing["mean_mli_um"])
    assert pd.isna(missing["field_balanced_mean_mli_um"])
    assert pd.isna(missing["sd_mli_um"])
    assert pd.isna(missing["field_sd_mli_um"])
    assert pd.isna(missing["field_sem_mli_um"])
    assert pd.isna(missing["sem_mli_um"])
    assert pd.isna(missing["field_balanced_mean_non_airspace_chord_um"])
    assert pd.isna(missing["field_sem_non_airspace_chord_um"])

    assert mixed["mli_measured_field_count"] == 1
    assert mixed["non_airspace_measured_field_count"] == 1
    assert np.isclose(mixed["mean_mli_um"], 12.0)
    assert np.isclose(mixed["field_balanced_mean_mli_um"], 12.0)
    assert np.isclose(mixed["sd_mli_um"], 0.0)
    assert np.isclose(mixed["field_sd_mli_um"], 0.0)
    assert np.isclose(mixed["field_sem_mli_um"], 0.0)
    assert np.isclose(mixed["sem_mli_um"], 0.0)
    assert np.isclose(mixed["mean_non_airspace_chord_um"], 7.0)
    assert np.isclose(mixed["sd_non_airspace_chord_um"], 0.0)
    assert np.isclose(mixed["field_sem_non_airspace_chord_um"], 0.0)


def test_indirect_equivalent_blank_when_accepted_edge_chords_lack_boundary_pairs():
    mask = np.array([[True, True, False, True, True, False]], dtype=bool)
    lines = generate_grid(mask.shape, orientation="horizontal", strategy="count", num_lines=1)
    chords = measure_phase_chords(
        mask,
        lines,
        phase="airspace",
        pixel_width_um=1.0,
        pixel_height_um=1.0,
        exclude_edge_touching=False,
    )

    summary = summarize_measurements(chords, "MLI")
    summary.update(
        summarize_phase_line_intercepts(
            mask,
            lines,
            phase="airspace",
            pixel_width_um=1.0,
            pixel_height_um=1.0,
            exclude_edge_touching=False,
            measure_type="MLI",
        )
    )

    assert summary["mli_chord_count"] == 2
    assert summary["mli_chord_count_raw"] == 2
    assert np.isclose(summary["mli_total_airspace_line_length_um_accepted"], 4.0)
    assert summary["mli_airspace_boundary_intersection_count_accepted"] == 3
    assert summary["mli_indirect_equivalent_mean_um"] is None
    assert summary["mli_direct_mean_um"] == summary["mli_mean_um"]


def test_orientation_balanced_indirect_blank_when_any_accepted_orientation_is_incomplete():
    mask = np.zeros((5, 5), dtype=bool)
    mask[2, 1:4] = True
    mask[0:2, 2] = True
    lines = generate_grid(mask.shape, orientation="both", strategy="count", num_lines=1)
    chords = measure_phase_chords(
        mask,
        lines,
        phase="airspace",
        pixel_width_um=1.0,
        pixel_height_um=1.0,
        exclude_edge_touching=False,
    )

    summary = summarize_measurements(chords, "MLI")
    summary.update(
        summarize_phase_line_intercepts(
            mask,
            lines,
            phase="airspace",
            pixel_width_um=1.0,
            pixel_height_um=1.0,
            exclude_edge_touching=False,
            measure_type="MLI",
        )
    )

    assert summary["mli_horizontal_count"] == 1
    assert summary["mli_vertical_count"] == 1
    assert np.isclose(summary["mli_indirect_equivalent_horizontal_mean_um"], 3.0)
    assert summary["mli_indirect_equivalent_vertical_mean_um"] is None
    assert summary["mli_indirect_equivalent_mean_um"] is None
    assert summary["mli_indirect_equivalent_orientation_balanced_mean_um"] is None



def test_component_connectivity_is_recorded_in_summaries_and_parameters(tmp_path):
    arr = np.zeros((4, 4, 3), dtype=np.uint8)
    arr[0, 0] = [255, 255, 255]
    arr[1, 1] = [255, 255, 255]
    arr[2, 2] = [255, 255, 255]
    img_path = tmp_path / "SlideConn_0001.png"
    Image.fromarray(arr).save(img_path)

    params_8 = AnalysisParams(
        pixel_width_um=1,
        pixel_height_um=1,
        threshold_method="otsu",
        airspace_component_connectivity=8,
    )
    params_4 = AnalysisParams(
        pixel_width_um=1,
        pixel_height_um=1,
        threshold_method="otsu",
        airspace_component_connectivity=4,
    )
    no_reference_8 = analyze_image(img_path, params_8)
    no_reference_4 = analyze_image(img_path, params_4)

    assert no_reference_8.summary["area_measurement_connectivity"] == airspace_component_connectivity_label()
    assert no_reference_8.summary["airspace_component_connectivity"] == 8
    assert no_reference_8.summary["airspace_component_count_all"] == 1
    assert no_reference_8.summary["final_airspace_area_pixels_mask"] == 0
    assert np.isclose(no_reference_8.summary["non_edge_airspace_fraction"], 0.0)
    assert no_reference_4.summary["area_measurement_connectivity"] == airspace_component_connectivity_label(4)
    assert no_reference_4.summary["airspace_component_connectivity"] == 4
    assert no_reference_4.summary["airspace_component_count_all"] == 3
    assert no_reference_4.summary["final_airspace_area_pixels_mask"] == 2
    assert np.isclose(no_reference_4.summary["non_edge_airspace_fraction"], 2 / 16)

    out = tmp_path / "connectivity_batch"
    result = process_files([img_path], out, params_4)
    field_row = result["field_summary"].iloc[0]
    slide_row = result["slide_summary"].iloc[0]
    assert field_row["area_measurement_connectivity"] == airspace_component_connectivity_label(4)
    assert slide_row["airspace_component_connectivity"] == 4
    assert slide_row["area_measurement_connectivity"] == airspace_component_connectivity_label(4)
    with open(out / "parameters.json", encoding="utf-8") as handle:
        parameters = json.load(handle)
    assert parameters["airspace_component_connectivity"] == 4
    assert parameters["area_measurement_connectivity"] == airspace_component_connectivity_label(4)
    assert "4-connected" in parameters["area_measurement_method_description"]


def test_random_grid_offsets_are_recorded_in_summaries_and_parameters(tmp_path):
    arr = _rgb_image((32, 32), np.s_[8:24, 8:24], background=(60, 60, 80))
    img_path = tmp_path / "SlideSeed_0001.png"
    Image.fromarray(arr).save(img_path)

    out = tmp_path / "random_grid"
    result = process_files(
        [img_path],
        out,
        AnalysisParams(
            pixel_width_um=1,
            pixel_height_um=1,
            num_lines=3,
            threshold_method="otsu",
            grid_random_offset=True,
            grid_random_seed=20260519,
        ),
    )

    field_row = result["field_summary"].iloc[0]
    assert bool(field_row["grid_random_offset"]) is True
    assert int(field_row["grid_random_seed"]) == 20260519
    assert int(field_row["grid_random_field_seed"]) >= 0
    assert np.isfinite(field_row["grid_horizontal_offset_px"])
    assert np.isfinite(field_row["grid_vertical_offset_px"])
    assert np.isclose(field_row["grid_horizontal_offset_um"], field_row["grid_horizontal_offset_px"])
    assert np.isclose(field_row["grid_vertical_offset_um"], field_row["grid_vertical_offset_px"])

    with open(out / "parameters.json", encoding="utf-8") as handle:
        parameters = json.load(handle)
    assert parameters["grid_random_offset"] is True
    assert parameters["grid_random_seed"] == 20260519

    with open(out / "SlideSeed_0001" / "summary.json", encoding="utf-8") as handle:
        image_summary = json.load(handle)
    assert image_summary["grid_random_seed"] == 20260519
    assert image_summary["grid_random_field_seed"] == int(field_row["grid_random_field_seed"])


def test_analyze_synthetic_field(tmp_path):
    arr = _rgb_image(
        (64, 64), np.s_[16:48, 8:24], np.s_[16:48, 40:56],
        background=(80, 60, 90), bright=(245, 245, 245)
    )
    img_path = tmp_path / "SlideA_0001.png"
    Image.fromarray(arr).save(img_path)
    result = analyze_image(
        img_path,
        AnalysisParams(pixel_width_um=1, pixel_height_um=1, num_lines=5, measure_non_airspace=True),
        tmp_path / "out",
    )
    assert result.summary["mli_chord_count"] > 0
    assert result.summary["mli_chord_count_accepted"] == result.summary["mli_chord_count"]
    assert result.summary["mli_chord_count_raw"] >= result.summary["mli_chord_count_include_edge"]
    assert result.summary["mli_chord_count_include_edge"] >= result.summary["mli_chord_count_accepted"]
    assert result.summary["mli_chord_count_min_length_excluded"] == (
        result.summary["mli_chord_count_raw"] - result.summary["mli_chord_count_include_edge"]
    )
    assert result.summary["mli_chord_count_edge_excluded"] == (
        result.summary["mli_chord_count_include_edge"] - result.summary["mli_chord_count_accepted"]
    )
    assert 0 <= result.summary["mli_chord_fraction_min_length_excluded"] <= 1
    assert 0 <= result.summary["mli_chord_fraction_edge_excluded"] <= 1
    assert result.summary["mli_mean_um_accepted"] == result.summary["mli_mean_um"]
    assert (
        result.summary["mli_orientation_balanced_mean_um_accepted"]
        == result.summary["mli_orientation_balanced_mean_um"]
    )
    assert result.summary["mli_direct_mean_um"] == result.summary["mli_mean_um"]
    assert result.summary["mli_pooled_chord_mean_um"] == result.summary["mli_mean_um"]
    assert result.summary["mli_direct_orientation_balanced_mean_um"] == result.summary["mli_orientation_balanced_mean_um"]
    assert result.summary["mli_orientation_balanced_direct_mean_um"] == result.summary["mli_orientation_balanced_mean_um"]
    assert result.summary["mli_direct_horizontal_vertical_mean_ratio"] is not None
    assert result.summary["mli_direct_horizontal_vertical_mean_delta_um"] is not None
    assert result.summary["mli_total_airspace_line_length_um_raw"] >= result.summary["mli_total_airspace_line_length_um_accepted"]
    assert result.summary["mli_airspace_boundary_intersection_count_raw"] >= result.summary["mli_airspace_boundary_intersection_count_accepted"]
    assert result.summary["mli_indirect_equivalent_orientation_balanced_mean_um"] is not None
    assert np.isclose(
        result.summary["mli_indirect_equivalent_orientation_balanced_mean_um"],
        result.summary["mli_direct_orientation_balanced_mean_um"],
    )
    assert result.summary["non_airspace_chord_count"] > 0
    assert result.summary["non_airspace_chord_count_accepted"] == result.summary["non_airspace_chord_count"]
    assert result.summary["non_airspace_chord_count_raw"] >= result.summary["non_airspace_chord_count_include_edge"]
    assert result.summary["non_airspace_chord_count_include_edge"] >= result.summary["non_airspace_chord_count_accepted"]
    assert result.summary["non_airspace_chord_count_min_length_excluded"] == (
        result.summary["non_airspace_chord_count_raw"] - result.summary["non_airspace_chord_count_include_edge"]
    )
    assert result.summary["non_airspace_chord_count_edge_excluded"] == (
        result.summary["non_airspace_chord_count_include_edge"] - result.summary["non_airspace_chord_count_accepted"]
    )
    assert 0 <= result.summary["non_airspace_chord_fraction_min_length_excluded"] <= 1
    assert 0 <= result.summary["non_airspace_chord_fraction_edge_excluded"] <= 1
    assert result.summary["non_airspace_mean_um_accepted"] == result.summary["non_airspace_mean_um"]
    assert (
        result.summary["non_airspace_orientation_balanced_mean_um_accepted"]
        == result.summary["non_airspace_orientation_balanced_mean_um"]
    )
    assert result.summary["non_airspace_direct_mean_um"] == result.summary["non_airspace_mean_um"]
    assert result.summary["non_airspace_pooled_chord_mean_um"] == result.summary["non_airspace_mean_um"]
    assert (
        result.summary["non_airspace_orientation_balanced_direct_mean_um"]
        == result.summary["non_airspace_orientation_balanced_mean_um"]
    )
    assert result.summary["mean_non_airspace_chord_um"] == result.summary["non_airspace_orientation_balanced_mean_um"]
    assert (
        result.summary["orientation_balanced_mean_non_airspace_chord_um"]
        == result.summary["non_airspace_orientation_balanced_mean_um"]
    )
    assert result.summary["pooled_mean_non_airspace_chord_um"] == result.summary["non_airspace_pooled_chord_mean_um"]
    assert "non_airspace_direct_horizontal_vertical_mean_ratio" in result.summary
    assert "non_airspace_direct_horizontal_vertical_mean_delta_um" in result.summary
    assert result.summary["total_airspace_area_um2_mask"] > 0
    assert result.summary["largest_airspace_component_area_um2"] > 0
    assert 0 < result.summary["largest_airspace_component_fraction_of_airspace"] <= 1
    assert (
        result.summary["non_edge_airspace_component_area_um2_mask"]
        == result.summary["final_airspace_area_um2_mask"]
    )
    assert result.summary["grid_random_offset"] is False
    assert result.summary["grid_random_seed"] is None
    assert result.summary["grid_horizontal_offset_px"] == 0.0
    assert result.summary["grid_vertical_offset_px"] == 0.0
    assert result.summary["calibration_source"] == UNRECORDED_CALIBRATION_SOURCE
    assert result.summary["calibration_is_default"] is False
    assert result.summary["calibration_warning"] == ""
    assert result.summary["pixel_area_um2"] == 1.0
    assert result.summary["field_width_um"] == 64.0
    assert result.summary["field_height_um"] == 64.0
    assert result.summary["input_image_type"] == "pre_cropped_lung_roi"
    assert result.summary["field_selection_stage"] == "upstream_pre_app"
    assert result.summary["field_selection_method"] == "not_recorded"
    assert result.summary["field_selected_by_app"] is False
    assert result.summary["field_included_in_summary"] is True
    assert result.summary["field_exclusion_flag"] is False
    assert result.summary["edge_exclusion_applied"] is True
    assert result.summary["short_chord_exclusion_applied"] is False
    assert result.overlay_path is not None
    assert result.overlay_path.exists()
    assert result.non_airspace_overlay_path is not None
    assert result.non_airspace_overlay_path.exists()
    assert result.final_contours_path is not None
    assert result.final_contours_path.exists()
    assert result.mask_overlay_path is not None
    assert result.mask_overlay_path.exists()
    assert result.mli_rejected_edge_overlay_path is not None
    assert result.mli_rejected_edge_overlay_path.exists()
    assert result.non_airspace_rejected_edge_overlay_path is not None
    assert result.non_airspace_rejected_edge_overlay_path.exists()
    image_dir = tmp_path / "out" / "SlideA_0001"
    assert (image_dir / "01_preprocessing" / "binary_airspace_mask.png").exists()
    assert not (image_dir / "binary_airspace_mask.tif").exists()
    assert (image_dir / "02_contours" / "final_airspace_contours_filled.png").exists()
    assert (image_dir / "03_mli" / "overlay_mli_combined.png").exists()
    assert (image_dir / "03_mli" / "grids" / "horizontal_grid_mli.png").exists()
    assert (image_dir / "04_non_airspace" / "overlay_non_airspace_combined.png").exists()
    assert (image_dir / "04_non_airspace" / "grids" / "horizontal_grid_non_airspace.png").exists()
    assert not (image_dir / "overlay_mli.tif").exists()
    assert not (image_dir / "overlay_non_airspace.tif").exists()
    assert (image_dir / "05_data" / "chords.csv").exists()
    assert (image_dir / "05_data" / "chords_audit.csv").exists()
    chord_export_columns = list(pd.read_csv(image_dir / "05_data" / "chords.csv", nrows=0).columns)
    assert chord_export_columns == [
        "Image",
        "Slide",
        "Field",
        "Measurement",
        "Orientation",
        "Line",
        "Line position (px)",
        "Length (µm)",
        "Length (px)",
        "Edge-touching?",
    ]
    audit_chord_columns = list(pd.read_csv(image_dir / "05_data" / "chords_audit.csv", nrows=0).columns)
    assert "start_x" in audit_chord_columns
    assert "length_um" in audit_chord_columns
    assert (image_dir / "06_qc" / "mask_overlay.png").exists()
    assert (image_dir / "06_qc" / "rejected_edge_chords_mli.png").exists()
    assert (image_dir / "06_qc" / "rejected_edge_chords_non_airspace.png").exists()
    assert not (image_dir / "chords.csv").exists()
    assert (image_dir / "qc_panel.png").exists()
    assert not (image_dir / "qc_panel.tif").exists()
    assert not (image_dir / "qc_panel_preview.png").exists()


def test_process_files_disambiguates_duplicate_input_stems_without_schema_drift(tmp_path):
    arr = _rgb_image((16, 16), np.s_[4:12, 4:12], background=(40, 40, 40))
    first_dir = tmp_path / "first"
    second_dir = tmp_path / "second"
    first_dir.mkdir()
    second_dir.mkdir()
    first_image = first_dir / "SlideDup_0001.png"
    second_image = second_dir / "SlideDup_0001.png"
    Image.fromarray(arr).save(first_image)
    arr_variant = arr.copy()
    arr_variant[6:10, 6:10] = [255, 255, 255]
    Image.fromarray(arr_variant).save(second_image)

    out = tmp_path / "duplicate_batch"
    result = process_files(
        [first_dir, second_dir],
        out,
        AnalysisParams(
            pixel_width_um=1,
            pixel_height_um=1,
            threshold_method="otsu",
            num_lines=2,
            exclude_edge_touching=False,
        ),
    )

    image_dirs = sorted(
        path.name for path in out.iterdir() if path.is_dir() and path.name.startswith("SlideDup")
    )
    assert image_dirs == ["SlideDup_0001", "SlideDup_0001__2"]
    for image_dir in image_dirs:
        per_image_dir = out / image_dir
        assert (per_image_dir / "03_mli" / "overlay_mli_combined.png").exists()
        assert (per_image_dir / "05_data" / "chords.csv").exists()
        assert (per_image_dir / "qc_panel.png").exists()

    assert len(result["field_summary"]) == 2
    assert len(result["slide_summary"]) == 1
    assert result["slide_summary"].iloc[0]["field_count"] == 2
    assert result["field_summary"]["filename"].tolist() == ["SlideDup_0001.png", "SlideDup_0001.png"]
    assert result["field_summary"]["slide_id"].tolist() == ["SlideDup", "SlideDup"]
    assert result["field_summary"]["field_id"].tolist() == ["0001", "0001"]
    assert "_input_index" not in result["field_summary"].columns
    assert "_input_path" not in result["field_summary"].columns
    ok_logs = result["processing_log"][result["processing_log"]["status"] == "ok"]
    assert ok_logs["filename"].tolist() == ["SlideDup_0001.png", "SlideDup_0001.png"]
    assert "_input_index" not in result["processing_log"].columns
    assert "_input_path" not in result["processing_log"].columns
    assert len(result["preview_paths"]) == 2
    assert len(set(result["preview_paths"])) == 2
    assert {path.parent.name for path in result["preview_paths"]} == set(image_dirs)
    assert all(path.exists() for path in result["preview_paths"])

    compact_chord_columns = list(pd.read_csv(out / "all_chords.csv", nrows=0).columns)
    assert compact_chord_columns[:3] == ["Image", "Slide", "Field"]
    audit_chord_columns = list(pd.read_csv(out / "audit" / "all_chords_audit.csv", nrows=0).columns)
    assert audit_chord_columns[:3] == ["filename", "slide_id", "field_id"]

    audit_workbook = load_workbook(out / "audit" / "lingappan_mli_audit.xlsx", read_only=True)
    field_headers = [
        cell for cell in next(audit_workbook["field_summary"].iter_rows(max_row=1, values_only=True))
    ]
    log_headers = [
        cell for cell in next(audit_workbook["processing_log"].iter_rows(max_row=1, values_only=True))
    ]
    assert field_headers[:3] == ["filename", "slide_id", "field_id"]
    assert "_input_index" not in field_headers
    assert "_input_path" not in field_headers
    assert "_input_index" not in log_headers
    assert "_input_path" not in log_headers



def _write_tiny_batch_image(tmp_path, name="SlideTiny_0001.png"):
    arr = _rgb_image((12, 12), np.s_[3:9, 3:9], background=(40, 40, 40))
    img_path = tmp_path / name
    Image.fromarray(arr).save(img_path)
    return img_path


EXPECTED_COMPACT_CHORD_COLUMNS = [
    "Image",
    "Slide",
    "Field",
    "Measurement",
    "Orientation",
    "Line",
    "Line position (px)",
    "Length (µm)",
    "Length (px)",
    "Edge-touching?",
]

EXPECTED_FIELD_WORKBOOK_COLUMNS = [
    "Image",
    "Slide",
    "Field",
    "MLI, orientation-balanced (µm)",
    "MLI pooled chord mean (µm)",
    "MLI chords (n)",
    "MLI horizontal mean (µm)",
    "MLI vertical mean (µm)",
    "MLI H/V ratio",
    "All airspace fraction",
    "Non-edge airspace fraction",
    "Non-edge airspace area (µm²)",
    "Largest component fraction",
    "Non-airspace chord, orientation-balanced (µm)",
    "Non-airspace pooled chord mean (µm)",
    "Non-airspace chords (n)",
    "Calibration warning",
    "Image warning",
]

EXPECTED_SLIDE_WORKBOOK_COLUMNS = [
    "Slide",
    "Fields (n)",
    "MLI measured fields (n)",
    "Mean MLI, field-balanced (µm)",
    "MLI SD across fields (µm)",
    "MLI SEM (µm)",
    "MLI chord-pooled mean (µm)",
    "Mean MLI chords/field",
    "Total MLI chords",
    "All airspace fraction",
    "Non-edge airspace fraction",
    "Non-edge airspace area (µm²)",
    "Largest component fraction",
    "Non-airspace measured fields (n)",
    "Mean non-airspace chord, field-balanced (µm)",
    "Non-airspace SD across fields (µm)",
    "Non-airspace SEM (µm)",
    "Non-airspace chord-pooled mean (µm)",
    "Mean non-airspace chords/field",
    "Total non-airspace chords",
    "Default calibration?",
    "Calibration source",
]

EXPECTED_FIELD_QC_WORKBOOK_COLUMNS = [
    "Image",
    "Slide",
    "Field",
    "MLI raw candidates",
    "MLI min-length excluded",
    "MLI edge excluded",
    "MLI accepted",
    "MLI min-length excluded fraction",
    "MLI edge excluded fraction",
    "MLI indirect-equivalent (µm)",
    "MLI accepted line length (µm)",
    "MLI accepted boundary intersections",
    "Non-airspace raw candidates",
    "Non-airspace min-length excluded",
    "Non-airspace edge excluded",
    "Non-airspace accepted",
    "Non-airspace min-length excluded fraction",
    "Non-airspace edge excluded fraction",
    "Airspace components",
    "Edge components excluded",
    "Largest component area (µm²)",
    "Largest component touches edge?",
    "Calibration warning",
    "Image warning",
]


def _workbook_headers(workbook, sheet_name: str) -> list[str]:
    return [cell for cell in next(workbook[sheet_name].iter_rows(max_row=1, values_only=True))]


def test_tiny_batch_export_manifest_and_schema_golden(tmp_path):
    img_path = _write_tiny_batch_image(tmp_path)
    out = tmp_path / "manifest_batch"

    process_files(
        [img_path],
        out,
        AnalysisParams(
            pixel_width_um=1,
            pixel_height_um=1,
            threshold_method="otsu",
            num_lines=2,
            measure_non_airspace=True,
        ),
    )

    assert {path.name for path in out.iterdir()} == {
        "README_results.md",
        "SlideTiny_0001",
        "all_chords.csv",
        "audit",
        "lingappan_mli_results.xlsx",
        "parameters.json",
    }
    image_dir = out / "SlideTiny_0001"
    assert {path.name for path in image_dir.iterdir()} == {
        "01_preprocessing",
        "02_contours",
        "03_mli",
        "04_non_airspace",
        "05_data",
        "06_qc",
        "qc_panel.png",
        "summary.json",
    }
    assert {path.name for path in (out / "audit").iterdir()} == {
        "all_chords_audit.csv",
        "lingappan_mli_audit.xlsx",
    }
    expected_relative_files = {
        "all_chords.csv",
        "README_results.md",
        "parameters.json",
        "lingappan_mli_results.xlsx",
        "audit/all_chords_audit.csv",
        "audit/lingappan_mli_audit.xlsx",
        "SlideTiny_0001/01_preprocessing/binary_airspace_mask.png",
        "SlideTiny_0001/01_preprocessing/grayscale_8bit.png",
        "SlideTiny_0001/02_contours/all_airspace_contours.png",
        "SlideTiny_0001/02_contours/final_airspace_contours.png",
        "SlideTiny_0001/02_contours/final_airspace_contours_filled.png",
        "SlideTiny_0001/03_mli/overlay_mli_combined.png",
        "SlideTiny_0001/04_non_airspace/overlay_non_airspace_combined.png",
        "SlideTiny_0001/05_data/chords.csv",
        "SlideTiny_0001/05_data/chords_audit.csv",
        "SlideTiny_0001/06_qc/mask_overlay.png",
        "SlideTiny_0001/06_qc/rejected_edge_chords_mli.png",
        "SlideTiny_0001/06_qc/rejected_edge_chords_non_airspace.png",
        "SlideTiny_0001/qc_panel.png",
        "SlideTiny_0001/summary.json",
    }
    actual_relative_files = {
        path.relative_to(out).as_posix()
        for path in out.rglob("*")
        if path.is_file()
    }
    assert expected_relative_files.issubset(actual_relative_files)

    assert list(pd.read_csv(out / "all_chords.csv", nrows=0).columns) == EXPECTED_COMPACT_CHORD_COLUMNS
    assert (
        list(pd.read_csv(image_dir / "05_data" / "chords.csv", nrows=0).columns)
        == EXPECTED_COMPACT_CHORD_COLUMNS
    )

    workbook = load_workbook(out / "lingappan_mli_results.xlsx", read_only=True)
    assert workbook.sheetnames == [
        "Slides",
        "Fields",
        "Field QC",
        "Run settings",
        "Run log",
        "Column dictionary",
        "Export notes",
    ]
    assert _workbook_headers(workbook, "Fields") == EXPECTED_FIELD_WORKBOOK_COLUMNS
    assert _workbook_headers(workbook, "Slides") == EXPECTED_SLIDE_WORKBOOK_COLUMNS
    assert _workbook_headers(workbook, "Field QC") == EXPECTED_FIELD_QC_WORKBOOK_COLUMNS

    audit_workbook = load_workbook(out / "audit" / "lingappan_mli_audit.xlsx", read_only=True)
    assert audit_workbook.sheetnames == ["field_summary", "slide_summary", "processing_log", "parameters"]

    readme = (out / "README_results.md").read_text(encoding="utf-8")
    for phrase in (
        "pre-cropped lung fields/ROIs",
        "lingappan_mli_results.xlsx",
        "audit/lingappan_mli_audit.xlsx",
        "mean length of accepted continuous airspace segments",
        "not septal wall-thickness measurements",
        "06_qc/",
    ):
        assert phrase in readme


def test_compact_workbook_failure_still_attempts_audit_workbook_and_logs_warning(tmp_path, monkeypatch):
    img_path = _write_tiny_batch_image(tmp_path, "SlideFailCompact_0001.png")
    out = tmp_path / "compact_failure"
    attempted: list[str] = []
    real_writer = analysis_module.write_results_workbook

    def fail_compact(path, sheets):
        attempted.append(path.name)
        if path.name == "lingappan_mli_results.xlsx":
            raise RuntimeError("simulated compact workbook failure")
        real_writer(path, sheets)

    monkeypatch.setattr(analysis_module, "write_results_workbook", fail_compact)

    result = process_files(
        [img_path],
        out,
        AnalysisParams(pixel_width_um=1, pixel_height_um=1, threshold_method="otsu", num_lines=2),
    )

    assert attempted == ["lingappan_mli_results.xlsx", "lingappan_mli_audit.xlsx"]
    assert not (out / "lingappan_mli_results.xlsx").exists()
    assert (out / "audit" / "lingappan_mli_audit.xlsx").exists()
    warning_rows = result["processing_log"][result["processing_log"]["status"] == "warning"]
    assert "lingappan_mli_results.xlsx" in warning_rows["filename"].tolist()
    assert warning_rows["message"].str.contains("simulated compact workbook failure", regex=False).any()

    audit_workbook = load_workbook(out / "audit" / "lingappan_mli_audit.xlsx", read_only=True)
    processing_log_rows = list(audit_workbook["processing_log"].iter_rows(values_only=True))
    assert any(row[0] == "lingappan_mli_results.xlsx" for row in processing_log_rows[1:])
    readme = (out / "README_results.md").read_text(encoding="utf-8")
    assert "lingappan_mli_results.xlsx" in readme
    persisted_log = pd.read_csv(out / "audit" / "processing_log.csv")
    assert "lingappan_mli_results.xlsx" in persisted_log["filename"].tolist()


def test_audit_workbook_failure_keeps_compact_workbook_and_persists_warning(tmp_path, monkeypatch):
    img_path = _write_tiny_batch_image(tmp_path, "SlideFailAudit_0001.png")
    out = tmp_path / "audit_failure"
    attempted: list[str] = []
    real_writer = analysis_module.write_results_workbook

    def fail_audit(path, sheets):
        attempted.append(path.name)
        if path.name == "lingappan_mli_audit.xlsx":
            raise RuntimeError("simulated audit workbook failure")
        real_writer(path, sheets)

    monkeypatch.setattr(analysis_module, "write_results_workbook", fail_audit)

    result = process_files(
        [img_path],
        out,
        AnalysisParams(pixel_width_um=1, pixel_height_um=1, threshold_method="otsu", num_lines=2),
    )

    assert attempted == ["lingappan_mli_results.xlsx", "lingappan_mli_audit.xlsx"]
    assert (out / "lingappan_mli_results.xlsx").exists()
    assert not (out / "audit" / "lingappan_mli_audit.xlsx").exists()
    warning_rows = result["processing_log"][result["processing_log"]["status"] == "warning"]
    assert "lingappan_mli_audit.xlsx" in warning_rows["filename"].tolist()
    assert warning_rows["message"].str.contains("simulated audit workbook failure", regex=False).any()

    readme = (out / "README_results.md").read_text(encoding="utf-8")
    assert "Export warnings" in readme
    assert "lingappan_mli_audit.xlsx" in readme
    persisted_log = pd.read_csv(out / "audit" / "processing_log.csv")
    assert "lingappan_mli_audit.xlsx" in persisted_log["filename"].tolist()




def test_batch_writes_formatted_workbook_and_summaries(tmp_path):
    arr = _rgb_image((64, 64), np.s_[12:52, 12:52], background=(70, 70, 90))
    img_path = tmp_path / "SlideB_0001.png"
    Image.fromarray(arr).save(img_path)
    out = tmp_path / "batch"
    params = AnalysisParams(
        pixel_width_um=1,
        pixel_height_um=1,
        measure_non_airspace=True,
        field_selection_method="systematic random ROI sampling",
        field_selection_notes="blinded operator selected fields before upload",
        field_exclusion_criteria="exclude large airways/vessels before upload",
    )
    result = process_files([img_path], out, params)
    assert len(result["field_summary"]) == 1
    field_row = result["field_summary"].iloc[0]
    slide_row = result["slide_summary"].iloc[0]
    assert field_row["input_image_type"] == "pre_cropped_lung_roi"
    assert field_row["field_selection_method"] == "systematic random ROI sampling"
    assert field_row["field_selection_notes"] == "blinded operator selected fields before upload"
    assert field_row["field_exclusion_criteria"] == "exclude large airways/vessels before upload"
    assert bool(field_row["field_exclusion_flag"]) is False
    assert field_row["area_measurement_method"] == AREA_MEASUREMENT_METHOD
    assert field_row["airspace_component_connectivity"] == 8
    assert field_row["area_measurement_connectivity"] == airspace_component_connectivity_label()
    assert field_row["area_column_semantics"] == AREA_COLUMN_SEMANTICS
    assert field_row["calibration_source"] == UNRECORDED_CALIBRATION_SOURCE
    assert bool(field_row["calibration_is_default"]) is False
    assert field_row["calibration_warning"] == ""
    assert field_row["pixel_area_um2"] == 1
    assert field_row["field_width_um"] == 64
    assert field_row["field_height_um"] == 64
    assert slide_row["calibration_source"] == UNRECORDED_CALIBRATION_SOURCE
    assert slide_row["pixel_area_um2"] == 1
    assert slide_row["mean_field_width_um"] == 64
    assert slide_row["mean_field_height_um"] == 64
    assert slide_row["mean_field_area_um2"] == 4096
    assert field_row["mli_direct_orientation_balanced_mean_um"] == field_row["mli_orientation_balanced_mean_um"]
    assert field_row["mli_orientation_balanced_direct_mean_um"] == field_row["mli_orientation_balanced_mean_um"]
    assert field_row["mli_pooled_chord_mean_um"] == field_row["mli_mean_um"]
    assert field_row["mean_non_airspace_chord_um"] == field_row["non_airspace_orientation_balanced_mean_um"]
    assert (
        field_row["orientation_balanced_mean_non_airspace_chord_um"]
        == field_row["non_airspace_orientation_balanced_mean_um"]
    )
    assert field_row["pooled_mean_non_airspace_chord_um"] == field_row["non_airspace_pooled_chord_mean_um"]
    assert (
        field_row["non_edge_airspace_component_area_um2_mask"]
        == field_row["final_airspace_area_um2_mask"]
    )
    assert slide_row["airspace_component_connectivity"] == 8
    assert slide_row["area_measurement_connectivity"] == airspace_component_connectivity_label()
    assert "mean_non_edge_airspace_fraction" in slide_row.index
    assert "mean_non_edge_airspace_component_area_um2_mask" in slide_row.index
    assert "mean_largest_airspace_component_area_um2" in slide_row.index
    assert "mean_largest_airspace_component_fraction_of_airspace" in slide_row.index
    assert np.isclose(
        slide_row["mean_non_edge_airspace_component_area_um2_mask"],
        slide_row["mean_final_airspace_area_um2_mask"],
    )
    assert "mean_mli_orientation_balanced_direct_um" in slide_row.index
    assert "mean_mli_pooled_chord_um" in slide_row.index
    assert "field_balanced_mean_mli_um" in slide_row.index
    assert "field_sd_mli_um" in slide_row.index
    assert "field_sem_mli_um" in slide_row.index
    assert "chord_pooled_mean_mli_um" in slide_row.index
    assert "mean_mli_chords_per_field" in slide_row.index
    assert np.isclose(slide_row["field_balanced_mean_mli_um"], slide_row["mean_mli_um"])
    assert np.isclose(slide_row["field_sem_mli_um"], 0.0)
    assert np.isclose(slide_row["mean_mli_chords_per_field"], slide_row["total_mli_chords"])
    assert "mean_mli_horizontal_vertical_mean_ratio" in slide_row.index
    assert "mean_mli_include_edge_um" in slide_row.index
    assert "mean_mli_edge_excluded_um" in slide_row.index
    assert "mean_mli_chord_fraction_min_length_excluded" in slide_row.index
    assert "mean_non_airspace_chord_um" in slide_row.index
    assert "field_balanced_mean_non_airspace_chord_um" in slide_row.index
    assert "field_sd_non_airspace_chord_um" in slide_row.index
    assert "field_sem_non_airspace_chord_um" in slide_row.index
    assert "chord_pooled_mean_non_airspace_chord_um" in slide_row.index
    assert "mean_non_airspace_chords_per_field" in slide_row.index
    assert "mean_orientation_balanced_non_airspace_chord_um" in slide_row.index
    assert "mean_pooled_non_airspace_chord_um" in slide_row.index
    assert "total_non_airspace_chords" in slide_row.index
    if slide_row["total_non_airspace_chords"] > 0:
        assert np.isclose(
            slide_row["field_balanced_mean_non_airspace_chord_um"],
            slide_row["mean_non_airspace_chord_um"],
        )
        assert np.isclose(slide_row["field_sem_non_airspace_chord_um"], 0.0)
    else:
        assert pd.isna(slide_row["field_balanced_mean_non_airspace_chord_um"])
        assert pd.isna(slide_row["field_sem_non_airspace_chord_um"])
    assert np.isclose(slide_row["mean_non_airspace_chords_per_field"], slide_row["total_non_airspace_chords"])
    assert "mean_non_airspace_include_edge_um" in slide_row.index
    assert "mean_non_airspace_edge_excluded_um" in slide_row.index
    assert "mean_non_airspace_chord_fraction_min_length_excluded" in slide_row.index
    assert slide_row["total_mli_chords_raw"] >= slide_row["total_mli_chords_include_edge"]
    assert slide_row["total_mli_chords_min_length_excluded"] == (
        slide_row["total_mli_chords_raw"] - slide_row["total_mli_chords_include_edge"]
    )
    assert slide_row["total_mli_chords_include_edge"] >= slide_row["total_mli_chords_accepted"]
    assert slide_row["total_mli_chords_edge_excluded"] == (
        slide_row["total_mli_chords_include_edge"] - slide_row["total_mli_chords_accepted"]
    )
    assert slide_row["total_mli_airspace_line_length_um_raw"] >= slide_row["total_mli_airspace_line_length_um_accepted"]
    log_row = result["processing_log"].iloc[0]
    assert log_row["field_selection_method"] == "systematic random ROI sampling"
    assert bool(log_row["field_exclusion_flag"]) is False
    assert log_row["area_measurement_method"] == AREA_MEASUREMENT_METHOD
    with open(out / "parameters.json", encoding="utf-8") as handle:
        parameters = json.load(handle)
    assert parameters["input_image_type"] == "pre_cropped_lung_roi"
    assert parameters["field_selection_stage"] == "upstream_pre_app"
    assert parameters["field_selection_method"] == "systematic random ROI sampling"
    assert parameters["field_exclusion_criteria"] == "exclude large airways/vessels before upload"
    assert parameters["area_measurement_method"] == AREA_MEASUREMENT_METHOD
    assert parameters["airspace_component_connectivity"] == 8
    assert parameters["area_measurement_connectivity"] == airspace_component_connectivity_label()
    assert parameters["area_column_semantics"] == AREA_COLUMN_SEMANTICS
    assert parameters["calibration_source"] == UNRECORDED_CALIBRATION_SOURCE
    assert parameters["calibration_is_default"] is False
    assert parameters["calibration_warning"] == ""
    assert parameters["pixel_area_um2"] == 1
    assert "geometric contour integration" in parameters["area_measurement_method_description"]
    assert not (out / "field_summary.csv").exists()
    assert not (out / "slide_summary.csv").exists()
    assert not (out / "processing_log.csv").exists()
    assert (out / "all_chords.csv").exists()
    assert (out / "audit" / "all_chords_audit.csv").exists()
    compact_chord_columns = list(pd.read_csv(out / "all_chords.csv", nrows=0).columns)
    assert compact_chord_columns == [
        "Image",
        "Slide",
        "Field",
        "Measurement",
        "Orientation",
        "Line",
        "Line position (px)",
        "Length (µm)",
        "Length (px)",
        "Edge-touching?",
    ]
    audit_chord_columns = list(pd.read_csv(out / "audit" / "all_chords_audit.csv", nrows=0).columns)
    assert "start_x" in audit_chord_columns
    assert "length_um" in audit_chord_columns
    workbook_path = out / "lingappan_mli_results.xlsx"
    assert workbook_path.exists()
    workbook = load_workbook(workbook_path, read_only=True)
    assert {
        "Slides",
        "Fields",
        "Field QC",
        "Run settings",
        "Run log",
        "Column dictionary",
        "Export notes",
    }.issubset(workbook.sheetnames)
    assert "field_audit" not in workbook.sheetnames
    field_headers = [cell for cell in next(workbook["Fields"].iter_rows(max_row=1, values_only=True))]
    assert "MLI, orientation-balanced (µm)" in field_headers
    assert "All airspace fraction" in field_headers
    assert "Non-edge airspace fraction" in field_headers
    assert "threshold_value" not in field_headers
    slide_headers = [cell for cell in next(workbook["Slides"].iter_rows(max_row=1, values_only=True))]
    assert "Mean MLI, field-balanced (µm)" in slide_headers
    assert "All airspace fraction" in slide_headers
    assert "Non-edge airspace fraction" in slide_headers
    assert "mean_mli_horizontal_vertical_mean_ratio" not in slide_headers
    qc_headers = [cell for cell in next(workbook["Field QC"].iter_rows(max_row=1, values_only=True))]
    assert "MLI edge excluded" in qc_headers
    assert "Image warning" in qc_headers
    dictionary_headers = [cell for cell in next(workbook["Column dictionary"].iter_rows(max_row=1, values_only=True))]
    assert {"Sheet", "Group", "Column", "Source column", "Units", "Meaning"}.issubset(dictionary_headers)
    assert "Audit location" not in dictionary_headers
    dictionary_rows = list(workbook["Column dictionary"].iter_rows(min_row=2, values_only=True))
    dictionary = {row[dictionary_headers.index("Column")]: row for row in dictionary_rows}
    mli_meaning = dictionary["Mean MLI, field-balanced (µm)"][dictionary_headers.index("Meaning")]
    assert "each field weighted equally" in mli_meaning
    non_edge_fraction_meaning = dictionary["Non-edge airspace fraction"][dictionary_headers.index("Meaning")]
    assert "fully internal/non-edge airspace components" in non_edge_fraction_meaning
    assert all(
        row[dictionary_headers.index("Meaning")] != "Primary display column derived from the source audit column."
        for row in dictionary_rows
    )
    audit_workbook_path = out / "audit" / "lingappan_mli_audit.xlsx"
    assert audit_workbook_path.exists()
    audit_workbook = load_workbook(audit_workbook_path, read_only=True)
    assert {"field_summary", "slide_summary", "processing_log", "parameters"}.issubset(audit_workbook.sheetnames)
    field_audit_headers = [cell for cell in next(audit_workbook["field_summary"].iter_rows(max_row=1, values_only=True))]
    assert "threshold_value" in field_audit_headers
    slide_audit_headers = [cell for cell in next(audit_workbook["slide_summary"].iter_rows(max_row=1, values_only=True))]
    assert "mean_mli_horizontal_vertical_mean_ratio" in slide_audit_headers
    assert len(field_headers) < len(field_audit_headers)
    assert len(slide_headers) < len(slide_audit_headers)
    export_notes = {
        row[0]: row[1]
        for row in workbook["Export notes"].iter_rows(min_row=2, values_only=True)
    }
    assert "geometric contour integration" in export_notes["Area measurement method"]
    assert "accepted continuous airspace segment" in export_notes["Chord measurement"]
    assert "Pixel size 1 × 1 µm/px" in export_notes["Calibration audit"]
    assert "not septal wall-thickness" in export_notes["Non-airspace chord caveat"]
    assert "mean_non_airspace_chord_um" in export_notes["Non-airspace chord caveat"]
    assert "minimum-length-excluded" in export_notes["Chord-filter sensitivity"]
    assert "fields as the experimental unit" in export_notes["Slide-level weighting and experimental unit"]
    assert "chord_pooled_mean_*" in export_notes["Slide-level weighting and experimental unit"]
    assert "rejected_edge_chords_mli.png" in export_notes["First-class QC overlays"]
    assert "chords_preview" not in workbook.sheetnames
